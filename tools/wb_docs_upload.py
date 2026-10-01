# -*- coding: utf-8 -*-
"""Загрузка деклараций/сертификатов соответствия на карточки WB из Excel-таблицы.

    .venv\\Scripts\\python.exe tools\\wb_docs_upload.py [таблица.xlsx] [--limit N] [--yes] [--second]

Таблица (первый лист, строка 1 — заголовки, как в «Декларации_шаблон.xlsx»):
    Артикул или nmID | Тип (декларация/сертификат) | Номер | Дата начала | Дата окончания | ТН ВЭД (необяз.)

Как работает (по образцу ранних скриптов для деклараций):
  * документ пишется в характеристики карточки: 15001135 «Номер декларации соответствия»,
    15001136 «Номер сертификата соответствия», 15001137 «Дата регистрации…», 15001138 «Дата окончания…»
    (+ ТН ВЭД в 15004139 и 15000001, если колонка заполнена) — WB сам сверяет номер с реестром;
  * карточка пересылается целиком (cards/update), меняются только эти поля;
  * НИКОГДА не трогаем карточки, где документ уже одобрен или на проверке
    (README: повторная отправка сбрасывает срок проверки, ТН ВЭД на одобренных снимает одобрение);
  * если на карточке уже есть одобренный документ ДРУГОГО типа — по умолчанию пропуск
    (даты у сертификата и декларации общие, можно испортить первый). Флаг --second разрешает;
  * сначала ПЛАН без отправки, потом вопрос «Отправить?»; --limit N — отправить только N (пилот);
  * через 90 с каждая карточка перечитывается: «встал / не встал», ошибки берутся из cards/error/list
    (WB отвечает 200 и молча не применяет — известная особенность API);
  * отчёт: data\\declarations\\upload_ДД_ММ_ЧЧММ.xlsx, журнал: data\\wb_declarations_ledger.csv.
"""
import argparse
import csv
import json
import re
import sys
import time
from datetime import date, datetime
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from config import settings  # noqa: E402
from wb_catalog_snapshot import classify, status_ru  # noqa: E402

BASE = "https://content-api.wildberries.ru"
CH_DECL, CH_CERT, CH_START, CH_END = 15001135, 15001136, 15001137, 15001138
CH_TNVED, CH_TNVED_OFF = 15004139, 15000001
DECL_DIR = ROOT / "data" / "declarations"
LEDGER = ROOT / "data" / "wb_declarations_ledger.csv"
H = {"Authorization": settings.wb_api_key, "Content-Type": "application/json"}
S = requests.Session()
_charcs_cache: dict[int, set] = {}


# ---------- WB ----------
def wb(method, path, **kw):
    for attempt in range(8):
        try:
            r = S.request(method, BASE + path, headers=H, timeout=60, **kw)
            if r.status_code == 429:
                time.sleep(6 * (attempt + 1))
                continue
            if r.status_code == 401:
                sys.exit("WB ответил 401: ключ WB_API_KEY в .env недействителен.")
            return r
        except requests.RequestException as e:
            print(f"  сеть: {e}; повтор", flush=True)
            time.sleep(5 * (attempt + 1))
    raise RuntimeError(f"WB не отвечает: {path}")


def find_card(key: str) -> dict | None:
    key = str(key).strip()
    body = {"settings": {"filter": {"textSearch": key, "withPhoto": -1}, "cursor": {"limit": 100}}}
    cards = wb("POST", "/content/v2/get/cards/list", json=body).json().get("cards") or []
    if key.isdigit():
        hit = next((c for c in cards if str(c["nmID"]) == key), None)
        if hit:
            return hit
    hits = [c for c in cards if (c.get("vendorCode") or "").strip() == key]
    return hits[0] if len(hits) == 1 else None


def subject_charcs(sid: int) -> set:
    if sid not in _charcs_cache:
        r = wb("GET", f"/content/v2/object/charcs/{sid}")
        _charcs_cache[sid] = {c.get("charcID") for c in (r.json().get("data") or [])}
    return _charcs_cache[sid]


def error_list() -> object:
    for method, kw in (("POST", {"json": {"cursor": {"limit": 100}, "order": {"ascending": False}}}), ("GET", {})):
        try:
            r = wb(method, "/content/v2/cards/error/list", **kw)
            if r.status_code == 200:
                return r.json()
        except Exception:
            pass
    return None


def errors_for(errs, vc: str, nm: int) -> str:
    found = []

    def walk(x):
        if isinstance(x, dict):
            if (x.get("vendorCode") or "").strip() == vc.strip() or x.get("nmID") == nm:
                found.extend(map(str, x.get("errors") or []))
            v = x.get(vc) or x.get(vc.strip())
            if isinstance(v, list):
                found.extend(map(str, v))
            for y in x.values():
                walk(y)
        elif isinstance(x, list):
            for y in x:
                walk(y)
    walk(errs)
    return "; ".join(dict.fromkeys(found))[:400]


# ---------- данные ----------
def norm_type(t) -> str | None:
    t = str(t or "").strip().lower()
    if t.startswith("декл") or t in ("д", "decl", "declaration"):
        return "decl"
    if t.startswith("серт") or t in ("с", "c", "cert", "certificate"):
        return "cert"
    return None


def norm_date(v) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s[:10], fmt).strftime("%d.%m.%Y")
        except ValueError:
            pass
    return None


def wb_number(num: str, kind: str) -> str:
    """Номер в виде, в котором WB находит его в реестре (README «ДЕКЛАРАЦИИ»):
    BY-декларации — «ЕАЭС № BY/112 …» (со «№» и пробелом), BY-сертификаты — без «№»;
    остальным добавляем «ЕАЭС », если его нет."""
    n = re.sub(r"\s+", " ", str(num).strip())
    body = re.sub(r"^ЕАЭС\s*(№|N)?\s*", "", n, flags=re.I)
    if body.upper().startswith("BY"):
        return f"ЕАЭС № {body}" if kind == "decl" else f"ЕАЭС {body}"
    if kind == "decl" and re.match(r"RU\s*Д", body, re.I):  # так WB хранит RU-декларации: «ЕАЭС N RU Д-…»
        return f"ЕАЭС N {body}"
    return n if n.upper().startswith("ЕАЭС") else f"ЕАЭС {n}"


def key_num(s: str) -> str:
    return re.sub(r"[^0-9A-ZА-Я]", "", str(s).upper().replace("N", "").replace("№", ""))


def read_table(path: Path) -> list[dict]:
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = list(r) + [None] * 6
        if not r[0] and not r[2]:
            continue
        rows.append({"row": i, "key": str(r[0]).strip() if r[0] is not None else "",
                     "type": norm_type(r[1]), "number": str(r[2] or "").strip(),
                     "start": norm_date(r[3]), "end": norm_date(r[4]),
                     "tnved": re.sub(r"\D", "", str(r[5] or ""))})
    return rows


# ---------- логика ----------
def plan_row(x: dict, allow_second: bool) -> dict:
    if not x["key"]:
        return {**x, "status": "ошибка: нет артикула"}
    if not x["type"]:
        return {**x, "status": "ошибка: тип не «декларация»/«сертификат»"}
    if not x["number"]:
        return {**x, "status": "ошибка: нет номера"}
    if not x["start"]:
        return {**x, "status": "ошибка: дата начала не распознана"}
    card = find_card(x["key"])
    if not card:
        return {**x, "status": "карточка не найдена (или артикул неоднозначен)"}
    _, st, _ = classify(card.get("documents"))
    out = {**x, "nmID": card["nmID"], "vendorCode": card.get("vendorCode", ""), "brand": card.get("brand", ""),
           "subject": card.get("subjectName", ""), "before": status_ru(st), "card": card}
    if st in ("verified", "pending", "valid_wait"):
        return {**out, "status": f"пропуск: {status_ru(st)}"}
    allowed = subject_charcs(card["subjectID"])
    field = CH_DECL if x["type"] == "decl" else CH_CERT
    if field not in allowed:
        return {**out, "status": "пропуск: у предмета нет поля для этого типа документа"}
    other = CH_CERT if x["type"] == "decl" else CH_DECL
    other_type = 1 if x["type"] == "decl" else 2  # 1 — сертификат, 2 — декларация
    has_other_valid = any(i.get("type") == other_type and (i.get("verdict") or {}).get("status") == 1
                          for i in (card.get("documents") or {}).get("items") or [])
    if has_other_valid and not allow_second:
        return {**out, "status": "пропуск: есть одобренный документ другого типа (даты общие; --second)"}
    num = wb_number(x["number"], x["type"])
    chars = [c for c in card.get("characteristics") or []
             if c.get("id") not in (field, CH_START, CH_END, CH_TNVED, CH_TNVED_OFF)
             and not (c.get("id") == other and not has_other_valid)]
    chars.append({"id": field, "value": [num]})
    chars.append({"id": CH_START, "value": [x["start"]]})
    if x["end"]:
        chars.append({"id": CH_END, "value": [x["end"]]})
    if x["tnved"]:
        for cid in (CH_TNVED, CH_TNVED_OFF):
            if cid in allowed:
                chars.append({"id": cid, "value": [x["tnved"]]})
            else:
                chars.extend(c for c in card.get("characteristics") or [] if c.get("id") == cid)
    else:
        chars.extend(c for c in card.get("characteristics") or [] if c.get("id") in (CH_TNVED, CH_TNVED_OFF))
    payload = {k: card.get(k) for k in ("nmID", "vendorCode", "brand", "title", "description", "dimensions", "sizes")}
    payload["characteristics"] = chars
    return {**out, "wb_number": num, "payload": payload, "status": "к отправке"}


def check_landed(p: dict) -> tuple[str, str]:
    card = find_card(str(p["nmID"]))
    if not card:
        return "не перечитана", ""
    items = (card.get("documents") or {}).get("items") or []
    mine = [i for i in items if key_num(i.get("number", "")) == key_num(p["wb_number"])]
    _, st, _ = classify(card.get("documents"))
    if not mine:
        return "НЕ ВСТАЛ", status_ru(st)
    v = mine[0].get("verdict")
    return ("встал, на проверке" if not v else "встал: " + ("одобрен" if v.get("status") == 1 else v.get("reason", "?"))), status_ru(st)


def write_report(rows, path):
    import openpyxl
    from openpyxl.styles import Font
    wb_ = openpyxl.Workbook()
    ws = wb_.active
    ws.title = "Загрузка"
    head = ["Строка", "Артикул/nmID", "nmID", "Бренд", "Предмет", "Тип", "Номер (как ушёл в WB)", "Дата начала",
            "Дата окончания", "ТН ВЭД", "Статус до", "Результат", "Статус после", "Ошибка WB"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["row"], r["key"], r.get("nmID"), r.get("brand"), r.get("subject"),
                   {"decl": "декларация", "cert": "сертификат"}.get(r["type"], ""), r.get("wb_number") or r["number"],
                   r["start"], r["end"], r["tnved"], r.get("before"), r["status"], r.get("after"), r.get("error")])
    for col, w in zip("ABCDEFGHIJKLMN", (8, 26, 12, 14, 24, 12, 36, 12, 12, 12, 26, 40, 26, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb_.save(path)


def _flush_keyboard():
    """Выкинуть всё, что набрано в окне ДО вопроса (случайные нажатия, пока шла проверка),
    чтобы на вопрос «Отправить?» отвечал только человек и только после того, как вопрос появился."""
    try:
        import msvcrt
        while msvcrt.kbhit():
            msvcrt.getwch()
    except ImportError:
        pass


def ask_count(total: int) -> int:
    """Двойное подтверждение живой отправки. 0 — отмена."""
    _flush_keyboard()
    ans = input(f"\nОтправить на WB {total} карточек? Это ЖИВЫЕ правки.\n"
                f"  ВСЕ — отправить все, число — только столько (пилот), Enter — отмена: ").strip()
    if ans.isdigit() and int(ans) > 0:
        n = min(int(ans), total)
    elif ans.upper() == "ВСЕ":
        n = total
    else:
        return 0
    _flush_keyboard()
    ok = input(f"Точно отправить {n} карточек? Напишите ДА: ").strip().upper()
    return n if ok == "ДА" else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("table", nargs="?", type=Path, default=DECL_DIR / "к_загрузке.xlsx")
    ap.add_argument("--limit", type=int, default=0, help="отправить не больше N карточек (пилот)")
    ap.add_argument("--yes", action="store_true", help="не спрашивать подтверждение")
    ap.add_argument("--second", action="store_true", help="разрешить второй документ при одобренном первом")
    ap.add_argument("--wait", type=int, default=90)
    args = ap.parse_args()
    if not args.table.exists():
        sys.exit(f"Нет файла {args.table}\nЗаполните «Декларации_шаблон.xlsx» и сохраните как {args.table.name} "
                 f"в {args.table.parent}")
    DECL_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%d_%m_%H%M")
    report = DECL_DIR / f"upload_{stamp}.xlsx"

    rows = read_table(args.table)
    print(f"Таблица: {args.table.name}, строк: {len(rows)}. Проверяю карточки на WB (только чтение)…\n", flush=True)
    planned = []
    for i, x in enumerate(rows, 1):
        try:
            planned.append(plan_row(x, args.second))
        except Exception as e:
            planned.append({**x, "status": f"ошибка: {e}"[:200]})
        if i % 25 == 0:
            print(f"  проверено {i}/{len(rows)}", flush=True)
        time.sleep(0.3)
    from collections import Counter
    cnt = Counter(p["status"] for p in planned)
    print("\nПЛАН:")
    for k, v in cnt.most_common():
        print(f"  {v:6}  {k}")
    todo = [p for p in planned if p["status"] == "к отправке"]
    if args.limit:
        todo = todo[:args.limit]
    write_report(planned, report)
    print(f"\nПлан сохранён: {report}")
    if not todo:
        print("Отправлять нечего.")
        return
    if not args.yes:
        n = ask_count(len(todo))
        if not n:
            print("Отменено, ничего не отправлено.")
            return
        todo = todo[:n]

    print(f"\nОтправляю {len(todo)} карточек…", flush=True)
    for i, p in enumerate(todo, 1):
        r = wb("POST", "/content/v2/cards/update", json=[p["payload"]])
        body = {}
        try:
            body = r.json()
        except ValueError:
            pass
        if r.status_code not in (200, 204) or body.get("error"):
            p["status"] = "WB отклонил запрос"
            p["error"] = (body.get("errorText") or str(body) or r.text)[:400]
        else:
            p["status"] = "отправлено"
        with LEDGER.open("a", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerow([datetime.now().isoformat(timespec="seconds"), p["nmID"], p["vendorCode"],
                                    p["type"], p["wb_number"], p["start"], p["end"], p["tnved"], p["status"]])
        if i % 10 == 0:
            print(f"  {i}/{len(todo)}", flush=True)
        time.sleep(0.7)

    sent = [p for p in todo if p["status"] == "отправлено"]
    if sent:
        print(f"\nЖду {args.wait} с, пока WB применит, и перечитываю…", flush=True)
        time.sleep(args.wait)
        errs = error_list()
        if errs is not None:
            (DECL_DIR / f"_errors_{stamp}.json").write_text(json.dumps(errs, ensure_ascii=False, indent=1),
                                                            encoding="utf-8")
        for p in sent:
            p["status"], p["after"] = check_landed(p)
            if p["status"] == "НЕ ВСТАЛ":
                p["error"] = errors_for(errs, p["vendorCode"], p["nmID"]) or "WB молча не применил — см. cards/error/list"
            time.sleep(0.4)
    write_report(planned, report)
    cnt = Counter(p["status"] for p in todo)
    print("\nИТОГ:")
    for k, v in cnt.most_common():
        print(f"  {v:6}  {k}")
    print(f"\nОтчёт: {report}\nВердикт WB приходит в течение часов–3 рабочих дней: проверять «Снимок_WB.bat».")


if __name__ == "__main__":
    main()
