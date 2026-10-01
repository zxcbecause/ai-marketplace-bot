# -*- coding: utf-8 -*-
"""Снимок всего каталога WB: у каждой карточки — статус документа соответствия. ТОЛЬКО ЧТЕНИЕ.

    .venv\\Scripts\\python.exe tools\\wb_catalog_snapshot.py [--out data\\declarations\\_wb_snapshot_ДД_ММ.json]

Пишет:
  --out JSON          — список карточек {nmID, vendorCode, brand, subject, title, doc, doc_status, verdict, documents}
  ..._summary.txt     — сводка по статусам (её же печатает в консоль)
  ..._samples.json    — по 3 примера сырого поля documents на каждый статус (для разбора)
  Снимок_WB_ДД_ММ.xlsx — таблица в REPORTS_DIR (по умолчанию data\\reports)

doc_status: verified / need_second / rejected: <причина> / valid_wait / pending / none_blocked / none
(verified/pending/none совместимы с decl_verdicts_check.py; остальное он считает отказом).
Карточки не меняются: используется только POST content/v2/get/cards/list.
Не запускать с таймаутом: 40 тыс. карточек грузятся 15–30 минут; прогресс сохраняется каждые 50 страниц.
"""
import argparse
import collections
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from config import settings  # noqa: E402

URL = "https://content-api.wildberries.ru/content/v2/get/cards/list"
PAGE = 100


REASONS_RU = {
    "document_dates_mismatch": "даты не совпадают с реестром",
    "document_not_found": "не найден в реестре",
    "document_inactive": "документ прекращён/недействителен",
    "document_missing": "документ не приложен",
    "document_type_mismatch": "неверный тип документа",
}
STATUS_RU = {
    "verified": "Одобрен",
    "need_second": "Нужен второй документ",
    "valid_wait": "Документ одобрен, ждёт итоговой проверки",
    "pending": "На проверке",
    "none": "Нет документа",
    "none_blocked": "Нет документа (WB требует)",
}


def _real(it):
    """Настоящий документ: тип 1/2 с цифрами в номере (отсекает «НЕ УКАЗАН», «НЕИЗВЕСТНО», «-»,
    которые бот ставил при создании карточек) или тип 3 (рег. удостоверение)."""
    if str(it.get("id", "")).startswith("00000000"):
        return False
    if it.get("type") == 3:
        return True
    return it.get("type") in (1, 2) and bool(re.search(r"\d", str(it.get("number") or "")))


def classify(docs):
    """Статус по полю documents карточки (разобрано по живому снимку 01.10.2026).
    items[].type: 1 — сертификат, 2 — декларация, 3 — рег. удостоверение, 6 — служебное (кол-во).
    overallVerdict.status: 1 — всё хорошо, 2 — проблема (reason documents_missing / documents_invalid).
    overallVerdict верим, только если он не старше последнего приложенного документа (README).
    documents_missing при одобренном документе = WB требует ещё один документ
    (как правило, пара «сертификат ТР ТС 004/020» + «декларация ТР ЕАЭС 037»)."""
    docs = docs or {}
    its = [i for i in docs.get("items") or [] if _real(i)]
    o = docs.get("overallVerdict") or {}
    if not its:
        st = "none_blocked" if o.get("status") == 2 else "none"
        return "", st, o
    its.sort(key=lambda i: i.get("createdAt", ""), reverse=True)
    number = "; ".join(i.get("number") or i.get("tradeName") or "" for i in its)
    verdicts = [i.get("verdict") or {} for i in its]
    bad = sorted({REASONS_RU.get(v.get("reason"), v.get("reason")) for v in verdicts if v.get("status") == 2})
    fresh = o and str(o.get("createdAt", ""))[:19] >= str(its[0].get("createdAt", ""))[:19]
    if fresh and o.get("status") == 1:
        st = "verified"
    elif bad:
        st = "rejected: " + ", ".join(bad)
    elif fresh and o.get("reason") == "documents_missing":
        st = "need_second"
    elif verdicts and all(v.get("status") == 1 for v in verdicts):
        st = "valid_wait"
    else:
        st = "pending"
    return number, st, o


def fetch_page(session, headers, cursor):
    body = {"settings": {"cursor": cursor, "filter": {"withPhoto": -1}}}
    for attempt in range(10):
        try:
            r = session.post(URL, headers=headers, json=body, timeout=60)
            if r.status_code == 429:
                time.sleep(6 * (attempt + 1))
                continue
            if r.status_code == 401:
                sys.exit("WB ответил 401: ключ WB_API_KEY в .env недействителен.")
            r.raise_for_status()
            return r.json()
        except requests.RequestException as e:
            print(f"  сеть: {e} — повтор через {5 * (attempt + 1)} с", flush=True)
            time.sleep(5 * (attempt + 1))
    sys.exit("WB не отвечает 10 раз подряд — снимок прерван (частичный файл сохранён).")


def save(rows, out):
    tmp = out.with_suffix(".tmp")
    tmp.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    tmp.replace(out)


def status_ru(st):
    return STATUS_RU.get(st) or ("Отклонён: " + st.split(": ", 1)[1] if st.startswith("rejected") else st)


def write_xlsx(rows, path):
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        return None
    bold = Font(bold=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка"
    cnt = collections.Counter(status_ru(r["doc_status"]) for r in rows)
    ws.append(["Статус документа", "Карточек"])
    for k, v in cnt.most_common():
        ws.append([k, v])
    ws.append(["Всего", len(rows)])
    for c in ws[1] + ws[ws.max_row]:
        c.font = bold
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 12

    order = ["Одобрен", "Нужен второй документ", "Отклонён", "Документ одобрен, ждёт итоговой проверки",
             "На проверке", "Нет документа (WB требует)", "Нет документа"]
    wsb = wb.create_sheet("По брендам")
    wsb.append(["Бренд", "Всего"] + order)
    by = collections.defaultdict(collections.Counter)
    for r in rows:
        k = status_ru(r["doc_status"])
        by[r["brand"] or "(без бренда)"][k.split(":")[0]] += 1
    for b, c in sorted(by.items(), key=lambda x: -sum(x[1].values())):
        wsb.append([b, sum(c.values())] + [c.get(k, 0) for k in order])
    for c in wsb[1]:
        c.font = bold
    wsb.column_dimensions["A"].width = 24
    wsb.freeze_panes = "B2"
    wsb.auto_filter.ref = wsb.dimensions

    ws2 = wb.create_sheet("Карточки")
    ws2.append(["nmID", "Артикул", "Бренд", "Предмет", "Название", "Документ(ы)", "Статус"])
    for c in ws2[1]:
        c.font = bold
    for r in rows:
        ws2.append([r["nmID"], r["vendorCode"], r["brand"], r["subject"], r["title"], r["doc"],
                    status_ru(r["doc_status"])])
    for col, w in zip("ABCDEFG", (12, 28, 16, 28, 50, 40, 40)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--from-json", type=Path, help="не качать заново, пересчитать статусы из готового снимка")
    ap.add_argument("--out", type=Path,
                    default=ROOT / "data" / "declarations" / f"_wb_snapshot_{datetime.now():%d_%m}.json")
    args = ap.parse_args()
    out = args.out if args.out.is_absolute() else ROOT / args.out
    out.parent.mkdir(parents=True, exist_ok=True)

    if args.from_json:
        rows = json.loads(args.from_json.read_text(encoding="utf-8"))
        for r in rows:
            r["doc"], r["doc_status"], r["verdict"] = classify(r["documents"])
        return report(rows, out, time.time())

    headers = {"Authorization": settings.wb_api_key, "Content-Type": "application/json"}
    session = requests.Session()
    cursor, rows, seen, page, t0 = {"limit": PAGE}, [], set(), 0, time.time()
    print(f"Снимок каталога WB → {out}\n(только чтение; 15–30 минут, окно не закрывать)\n", flush=True)
    while True:
        j = fetch_page(session, headers, cursor)
        cards = j.get("cards") or []
        for c in cards:
            if c["nmID"] in seen:
                continue
            seen.add(c["nmID"])
            docs = c.get("documents")
            number, st, verdict = classify(docs)
            rows.append({"nmID": c["nmID"], "vendorCode": c.get("vendorCode", ""), "brand": c.get("brand", ""),
                         "subject": c.get("subjectName", ""), "title": c.get("title", ""),
                         "doc": number, "doc_status": st, "verdict": verdict, "documents": docs})
        page += 1
        if page % 10 == 0:
            print(f"  карточек: {len(rows)}  ({time.time() - t0:.0f} с)", flush=True)
        if page % 50 == 0:
            save(rows, out)
        cur = j.get("cursor") or {}
        if len(cards) < PAGE or not cur.get("nmID"):
            break
        cursor = {"limit": PAGE, "updatedAt": cur["updatedAt"], "nmID": cur["nmID"]}
        time.sleep(0.6)
    save(rows, out)
    report(rows, out, t0)


def report(rows, out, t0):
    save(rows, out)
    cnt = collections.Counter(status_ru(r["doc_status"]) for r in rows)
    reasons = collections.Counter(f"{(r['verdict'] or {}).get('status')} / {(r['verdict'] or {}).get('reason')}"
                                  for r in rows)
    lines = [f"Снимок WB {datetime.now():%d.%m.%Y %H:%M}: карточек {len(rows)}", "", "Статус документа:"]
    lines += [f"  {k:<40} {v}" for k, v in cnt.most_common()]
    lines += ["", "Сырые пары status / reason от WB:"] + [f"  {v:6}  {k}" for k, v in reasons.most_common(40)]
    summary = "\n".join(lines)
    print("\n" + summary)
    out.with_name(out.stem + "_summary.txt").write_text(summary, encoding="utf-8")

    samples = collections.defaultdict(list)
    for r in rows:
        if len(samples[r["doc_status"]]) < 3:
            samples[r["doc_status"]].append({"nmID": r["nmID"], "vendorCode": r["vendorCode"],
                                             "documents": r["documents"]})
    out.with_name(out.stem + "_samples.json").write_text(json.dumps(samples, ensure_ascii=False, indent=1),
                                                         encoding="utf-8")
    xlsx = write_xlsx(rows, Path(os.getenv("REPORTS_DIR", str(ROOT / "data" / "reports"))) / f"Снимок_WB_{datetime.now():%d_%m}.xlsx")
    print(f"\nГотово за {(time.time() - t0) / 60:.1f} мин. Файлы:\n  {out}"
          + (f"\n  {xlsx}" if xlsx else ""))


if __name__ == "__main__":
    main()
