"""
Прайс-лист продавца (products_export.xlsx) — поиск цены по бренду+объёму
для категории SSD (find_price), либо по бренду+токенам модели для любой
другой категории (find_price_by_model). Колонка "Цена со скидкой для
покупателя ₸" — то, что указываем в карточке Ozon.

Сопоставление по бренду+объёму, а не по артикулу — внутренний "Артикул" в
прайсе (число вида 52645) не совпадает с артикулом продавца/SKU
производителя, который мы используем в Ozon-карточках.
"""
import logging
import os
import re
from pathlib import Path

import openpyxl

log = logging.getLogger(__name__)

PRICE_FILE = Path(os.getenv("PRICE_LIST_FILE", "./data/price_list.xlsx"))
SHEET = "Worksheet"
NAME_COL = 3       # C: Наименование товара
CATEGORY_COL = 5   # E: Категория товара
PRICE_COL = 25     # Y: Цена со скидкой для покупателя ₸

_CAPACITY_RE = re.compile(r'(\d+(?:[.,]\d+)?)\s*(TB|GB|Тб|ТБ|Гб|ГБ)', re.IGNORECASE)

_cache: list[tuple[str, str, float]] | None = None  # (наименование.lower(), категория.lower(), цена)


def _load() -> list[tuple[str, str, float]]:
    global _cache
    if _cache is not None:
        return _cache
    wb = openpyxl.load_workbook(PRICE_FILE, data_only=True)
    ws = wb[SHEET]
    rows: list[tuple[str, str, float]] = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=NAME_COL).value
        category = ws.cell(row=r, column=CATEGORY_COL).value
        price = ws.cell(row=r, column=PRICE_COL).value
        if not name or not price:
            continue
        try:
            price_f = float(price)
        except (TypeError, ValueError):
            continue
        if price_f <= 0:
            continue
        rows.append((str(name).lower(), str(category or "").lower(), price_f))
    _cache = rows
    log.info(f"Прайс-лист загружен: {len(rows)} позиций с ценой")
    return rows


def _capacity_gb(text: str) -> float | None:
    m = _CAPACITY_RE.search(text)
    if not m:
        return None
    num = float(m.group(1).replace(",", "."))
    unit = m.group(2).lower()
    if unit in ("tb", "тб"):
        num *= 1000
    return num


def find_price(brand: str, memory: str, model: str = "", category_hint: str = "ssd") -> int | None:
    """Ищет цену по бренду+объёму(+модели, если есть) среди позиций категории
    (по умолчанию — SSD/накопители). Возвращает округлённую "Цену со скидкой
    для покупателя ₸" или None, если совпадения не нашлось — в этом случае
    цену лучше оставить пустой, чем угадывать.
    model — доп. фильтр (например «SA500»), отсекает другие линейки того же
    бренда с тем же объёмом (Netac SA500 / N600S / NV2000 — разные модели)."""
    if not brand or not memory:
        return None
    target_gb = _capacity_gb(memory)
    if target_gb is None:
        return None

    brand_low = brand.strip().lower()
    # Первое "слово" модели обычно сама модельная линейка (SA500, N930E) —
    # длинные полные SKU с дефисами/суффиксами не совпадут с прайсом буквально.
    model_low = model.strip().lower().split()[0] if model.strip() else ""
    # Сравниваем по "хвосту" без знаков — у прайса и SKU могут различаться
    # префиксы серии (GIGABYTE "GP-GSTFS31480GNTD" в наших данных vs
    # "GSTFS31480GNTD" в прайсе), а конец кода обычно уникален и совпадает.
    model_key = re.sub(r"[^a-z0-9]", "", model_low)
    model_suffix = model_key[-8:] if len(model_key) > 8 else model_key
    rows = _load()

    def _search(require_model: bool) -> list[float]:
        result = []
        for name_low, category_low, price in rows:
            if category_hint not in category_low:
                continue
            if brand_low not in name_low:
                continue
            if require_model and model_suffix and model_suffix not in re.sub(r"[^a-z0-9]", "", name_low):
                continue
            gb = _capacity_gb(name_low)
            if gb is None or gb != target_gb:
                continue
            result.append(price)
        return result

    if model_low:
        # Модель известна — ищем строго бренд+модель+объём. Без отката на
        # "бренд+объём без модели": в прайсе у бренда часто несколько разных
        # линеек той же ёмкости (SA500/N600S/NV2000 у Netac) — взять цену
        # ЧУЖОЙ линейки хуже, чем оставить поле пустым.
        candidates = _search(require_model=True)
    else:
        candidates = _search(require_model=False)

    if not candidates:
        log.info(f"Цена не найдена в прайсе: бренд={brand!r} объём={memory!r} модель={model!r}")
        return None
    # Несколько совпадений (разные интерфейсы/линейки под тем же брендом+объёмом) —
    # берём минимальную, не угадывая конкретную модель/интерфейс.
    price = round(min(candidates))
    log.info(f"Цена из прайса: бренд={brand!r} объём={memory!r} модель={model!r} -> {price}")
    return price


def find_price_by_model(brand: str, model: str, category_hint: str = "") -> int | None:
    """Ищет цену по бренду + токенам модели без ёмкости/объёма.
    Подходит для акустики, наушников и других товаров, где нет GB/TB.
    category_hint — необязательный фильтр подстроки по колонке категории
    в прайсе; если пустой — фильтруем только по бренду+модели."""
    if not brand or not model:
        return None
    brand_low = brand.strip().lower()
    model_low = model.strip().lower()
    model_tokens = [t for t in re.split(r"[\s\-/]+", model_low) if len(t) >= 2]
    if not model_tokens:
        return None

    rows = _load()
    # (кол-во совпавших токенов, цена) — берём позицию с наибольшим совпадением
    best_count = 0
    best_price: float | None = None

    for name_low, category_low, price in rows:
        if category_hint and category_hint not in category_low:
            continue
        if brand_low not in name_low:
            continue
        match_count = sum(1 for t in model_tokens if t in name_low)
        if match_count == 0:
            continue
        if match_count > best_count:
            best_count = match_count
            best_price = price

    if best_price is None:
        log.info(f"Цена не найдена в прайсе (by_model): бренд={brand!r} модель={model!r}")
        return None
    price_int = round(best_price)
    log.info(
        f"Цена из прайса (by_model): бренд={brand!r} модель={model!r} "
        f"-> {price_int} (совпало {best_count} токенов)"
    )
    return price_int
